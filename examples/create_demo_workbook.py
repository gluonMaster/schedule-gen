"""Write fictional planning input. Run from any directory; existing output is kept."""
from pathlib import Path

from openpyxl import Workbook


def main():
    destination = Path(__file__).resolve().parent / "generated" / "demo_planning.xlsx"
    if destination.exists():
        raise SystemExit(f"Example already exists: {destination}")
    destination.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = "Plannung"
    labels = [
        "Subject", "Group", "Teacher", "Main room", "Alternative room 1",
        "Alternative room 2", "Alternative room 3", "Building", "Duration (minutes)",
        "Day", "Start time", "End of time window (optional)", "Pause before", "Pause after",
    ]
    lessons = [
        ["Demo mathematics", "Demo group A", "Demo teacher A", "V1.01", None, None,
         None, "Villa", 45, "Mo", "10:00", None, 0, 0],
        ["Demo science", "Demo group B", "Demo teacher B", "V1.02", None, None,
         None, "Villa", 45, "Mo", "11:00", None, 0, 0],
    ]
    # Section zero starts at row 2; columns C/D are reserved for linked lessons.
    for index, lesson in enumerate(lessons):
        first_row = 2 + index * 14
        for offset, (label, value) in enumerate(zip(labels, lesson)):
            sheet.cell(first_row + offset, 1, label)
            sheet.cell(first_row + offset, 2, value)
    sheet.column_dimensions["A"].width = 32
    sheet.column_dimensions["B"].width = 24
    book.save(destination)
    book.close()
    print(destination)


if __name__ == "__main__":
    main()
