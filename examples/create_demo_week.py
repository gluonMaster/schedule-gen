"""Write a fictional week of 24 lessons. Run from any directory; existing output is kept.

The rooms, groups and teachers are invented. Most lessons only have a time window,
so the optimizer has to choose the start times and resolve teacher and room conflicts.
"""
from pathlib import Path

from openpyxl import Workbook

LABELS = [
    "Subject", "Group", "Teacher", "Main room", "Alternative room 1",
    "Alternative room 2", "Alternative room 3", "Building", "Duration (minutes)",
    "Day", "Start time", "End of time window (optional)", "Pause before", "Pause after",
]

# subject, group, teacher, main room, alternative room, building, minutes, day, window start, window end
LESSONS = [
    ("Mathematik", "Mathe 5", "L. Becker", "V1.01", "V1.02", "Villa", 60, "Mo", "14:00", "18:00"),
    ("Mathematik", "Mathe 7", "L. Becker", "V1.01", "V1.02", "Villa", 60, "Mo", "14:00", "18:00"),
    ("Deutsch", "Deutsch 3", "M. Wagner", "V1.02", None, "Villa", 45, "Mo", "14:30", "17:30"),
    ("Englisch", "Englisch A1", "S. Fischer", "V1.03", None, "Villa", 45, "Mo", "15:00", "18:00"),
    ("Schach", "Schach Anfänger", "T. Schulz", "P1", None, "Haus am Park", 90, "Mo", "15:00", None),
    ("Kunst", "Kunst Kids", "K. Neumann", "P2", None, "Haus am Park", 90, "Di", "14:30", "18:00"),
    ("Mathematik", "Mathe 6", "L. Becker", "V1.01", None, "Villa", 60, "Di", "14:00", "18:00"),
    ("Englisch", "Englisch A2", "S. Fischer", "V1.03", "V1.02", "Villa", 60, "Di", "15:00", "18:30"),
    ("Gitarre", "Gitarre Gruppe", "J. Wolf", "V1.02", None, "Villa", 45, "Di", "16:00", "19:00"),
    ("Robotik", "Robotik 10+", "T. Schulz", "P1", None, "Haus am Park", 90, "Di", "16:00", None),
    ("Deutsch", "Deutsch 4", "M. Wagner", "V1.02", "V1.03", "Villa", 45, "Mi", "14:00", "17:00"),
    ("Theater", "Theater-AG", "K. Neumann", "P2", None, "Haus am Park", 90, "Mi", "15:00", "19:00"),
    ("Mathematik", "Mathe 8", "L. Becker", "V1.01", None, "Villa", 90, "Mi", "15:00", "19:00"),
    ("Englisch", "Englisch B1", "S. Fischer", "V1.03", None, "Villa", 60, "Mi", "15:30", "19:00"),
    ("Schach", "Schach Fortgeschr.", "T. Schulz", "P1", None, "Haus am Park", 90, "Do", "15:00", "19:00"),
    ("Kunst", "Malen & Zeichnen", "K. Neumann", "P2", None, "Haus am Park", 60, "Do", "14:30", "18:00"),
    ("Deutsch", "Deutsch 5", "M. Wagner", "V1.02", None, "Villa", 60, "Do", "14:00", "18:00"),
    ("Gitarre", "Gitarre Einsteiger", "J. Wolf", "V1.03", None, "Villa", 45, "Do", "16:00", "19:00"),
    ("Mathematik", "Mathe 9", "L. Becker", "V1.01", None, "Villa", 90, "Fr", "14:00", "18:00"),
    ("Tanz", "Tanz 6–8", "J. Wolf", "P2", None, "Haus am Park", 60, "Fr", "15:00", "18:00"),
    ("Englisch", "Englisch A1+", "S. Fischer", "V1.03", None, "Villa", 45, "Fr", "14:00", "17:00"),
    ("Robotik", "Robotik Kids", "T. Schulz", "P1", None, "Haus am Park", 90, "Sa", "10:00", "14:00"),
    ("Kunst", "Familienatelier", "K. Neumann", "P2", None, "Haus am Park", 120, "Sa", "10:00", "14:00"),
    ("Schach", "Schach-Turnier", "T. Schulz", "P1", "P2", "Haus am Park", 120, "Sa", "11:00", "16:00"),
]


def main():
    destination = Path(__file__).resolve().parent / "generated" / "demo_week.xlsx"
    if destination.exists():
        raise SystemExit(f"Example already exists: {destination}")
    destination.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = "Plannung"
    # One 14-row section per lesson, values in column B (see create_demo_workbook.py).
    for index, (subject, group, teacher, room, alternative, building, minutes, day, start, end) in enumerate(LESSONS):
        values = [subject, group, teacher, room, alternative, None, None, building,
                  minutes, day, start, end, 0, 0]
        first_row = 2 + index * 14
        for offset, (label, value) in enumerate(zip(LABELS, values)):
            sheet.cell(first_row + offset, 1, label)
            sheet.cell(first_row + offset, 2, value)
    sheet.column_dimensions["A"].width = 32
    sheet.column_dimensions["B"].width = 24
    book.save(destination)
    book.close()
    print(destination)


if __name__ == "__main__":
    main()
