"""Phase 3: actual Python XLSX boundaries; VBA transfer is inspected, not executed."""
import ast
import importlib
import json
import sys
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from gear_xls import excel_exporter, excel_parser, integration
from gear_xls.schedule_exchange import (
    RECORD_COLUMNS, ScheduleExchangeError, read_sync_metadata, write_sync_metadata,
)
from gear_xls.schedule_structure import build_schedule_structure
from gear_xls.services.schedule_pipeline import (
    SchedulePipeline, SchedulePipelineError, collect_individual_blocks_from_buildings,
    strip_non_group_activity_blocks_from_html,
)
from reader import ScheduleReader
from output_utils import export_to_excel, SCHEDULE_COLUMNS
from test_rental_model_editor import isolated_editor, login

SYNC = dict(format_version=1, source_base_revision='base-old', source_individual_revision='', snapshot_scope='full')


def records():
    common = dict(subject='Vermietung', teacher='Kontakt & Verein', students='Organisation',
                  building='Villa', room='1.01', day='Mo', start_time='10:00', end_time='11:00',
                  duration=60, lesson_type='rental', rental_dates=[], color='#abcdef')
    return [
        {**common, 'block_id': '00017', 'teacher': '', 'students': '', 'notes': 'Сохранить: ä & <>'},
        {**common, 'block_id': 'room-2', 'room': '1.02', 'teacher': '', 'students': 'Verein'},
        {**common, 'block_id': 'dated', 'room': '1.03', 'subject': 'Verein-Treffen', 'rental_dates': ['2026-09-28', '2026-10-05', '2026-10-12']},
        {**common, 'block_id': 'past', 'room': '1.04', 'rental_dates': ['2026-09-28']},
        {**common, 'block_id': 'trial', 'room': '1.05', 'subject': 'Trial', 'lesson_type': 'trial',
         'rental_dates': [], 'trial_dates': ['2026-10-05'], 'students': '', 'teacher': 'Trial Teacher'},
        {**common, 'block_id': '', 'room': '1.06', 'subject': 'Math', 'lesson_type': 'group',
         'students': '1A', 'teacher': 'Teacher'},
    ]


def write_planning_boundary(schedule_path, planning_path):
    """Fixture with the cells VBA must produce; deliberately not a VBA execution claim."""
    source = openpyxl.load_workbook(schedule_path)
    target = openpyxl.Workbook()
    planning = target.active
    planning.title = 'Plannung'
    metadata = target.create_sheet('__service_metadata')
    metadata.append(['section_index', 'column_letter', 'lesson_type', 'trial_dates_json', *RECORD_COLUMNS])
    for index, row in enumerate(source['Schedule'].iter_rows(min_row=2, values_only=True)):
        first = 2 + index * 14
        for offset, value in {0: row[0], 1: row[1], 2: row[2], 3: 'V' + row[3],
                              7: row[4], 8: row[8], 9: row[5], 10: row[6]}.items():
            planning.cell(first + offset, 2, value)
        if row[9] == 'rental':
            planning.cell(first + 11, 2, row[7])
        metadata.append([index, 'B', row[9], row[10], *row[11:]])
    write_sync_metadata(target, read_sync_metadata(source))
    target.save(planning_path)
    source.close()
    target.close()


def assert_activities(activities):
    assert len(activities) == 6
    by_id = {row['block_id']: row for row in activities.values()}
    for expected in records():
        row = by_id[expected['block_id']]
        for key in ('subject', 'lesson_type', 'teacher', 'students', 'rental_dates', 'start_time', 'end_time'):
            assert row[key] == expected[key]
        assert row['room_display'] == expected['room']
        assert row['building'] == expected['building']
    assert by_id['trial']['trial_dates'] == ['2026-10-05']
    assert by_id['00017']['block_metadata']['notes'] == 'Сохранить: ä & <>'


def test_python_workbook_boundaries_and_editor_output(tmp_path, monkeypatch):
    source, planning, optimized = [tmp_path / name for name in ('Schedule.xlsx', 'newpref.xlsx', 'optimized.xlsx')]
    assert excel_exporter.create_excel_from_html_data(records(), source, sync_metadata=SYNC) == source
    book = openpyxl.load_workbook(source)
    assert [cell.value for cell in book['Schedule'][1]][:11] == [
        'Занятие', 'Группа', 'Преподаватель', 'Кабинет', 'Здание', 'День',
        'Начало', 'Конец', 'Продолжительность', 'Тип занятия', 'Даты (JSON)']
    assert read_sync_metadata(book) == SYNC
    assert book['__schedule_sync'].sheet_state == 'veryHidden'
    book.close()
    activities = excel_parser.parse_schedule(source)
    assert_activities(activities)
    assert activities.sync_metadata == SYNC

    write_planning_boundary(source, planning)
    reader = ScheduleReader(planning)
    classes = reader.read_excel()
    assert reader.sync_metadata == SYNC
    assert [cls.block_id for cls in classes] == [row['block_id'] for row in records()]
    for cls, expected in zip(classes, records()):
        assert (cls.lesson_type, cls.rental_dates, cls.trial_dates) == (
            expected['lesson_type'], expected['rental_dates'], expected.get('trial_dates', []))
        assert (cls.teacher, cls.group) == (expected['teacher'], expected['students'])
        assert cls.main_room == 'V' + expected['room']
    assert classes[0].block_metadata['notes'] == 'Сохранить: ä & <>'

    # Test serialization of an output fixture independently of the unavailable solver.
    solution = [dict(row, group=row['students'], pause_before=0, pause_after=0) for row in records()]
    serializer = SimpleNamespace(solution=solution, sync_metadata=reader.sync_metadata, teachers=[], groups=[], rooms=[])
    assert export_to_excel(serializer, optimized)
    book = openpyxl.load_workbook(optimized)
    assert [cell.value for cell in book['Schedule'][1]][:11] == SCHEDULE_COLUMNS[:11]
    assert read_sync_metadata(book) == SYNC
    book.close()
    assert_activities(excel_parser.parse_schedule(optimized))

    html_dir = tmp_path / 'html'
    html_dir.mkdir()
    result = SchedulePipeline().process_excel_to_outputs(str(optimized), {'html': str(html_dir)}, spiski_data={})
    assert result['sync_metadata'] == SYNC
    assert len(result['individual_blocks']) == 5
    blocks = {row['id']: row for row in result['individual_blocks']}
    assert set(blocks) == {'00017', 'room-2', 'dated', 'past', 'trial'}
    assert blocks['past']['rental_dates'] == ['2026-09-28']
    assert blocks['00017']['notes'] == 'Сохранить: ä & <>'
    html = Path(result['html_file']).read_text(encoding='utf-8')
    assert "data-lesson-type='rental'" not in html
    assert "data-lesson-type='trial'" not in html
    assert "data-lesson-type='group'" in html
    assert 'block_metadata_json' in html  # current inline exporter was delivered
    monkeypatch.setattr(integration, 'get_schedule_state_dir', lambda: str(tmp_path / 'state'))
    # Phase 4 applies a snapshot only to the state it was exported from.
    (tmp_path / 'state').mkdir()
    (tmp_path / 'state/base_schedule.json').write_text(json.dumps({'published_at': 'base-old', 'blocks': []}), encoding='utf-8')
    (tmp_path / 'state/individual_lessons.json').write_text(json.dumps({'last_modified': None, 'blocks': [
        {'id': row['id']} for row in result['individual_blocks']]}), encoding='utf-8')
    integration.reset_web_editor_state(result['individual_blocks'], sync_metadata=result['sync_metadata'])
    state = json.loads((tmp_path / 'state/individual_lessons.json').read_text(encoding='utf-8'))
    assert state['blocks'] == result['individual_blocks']


@pytest.mark.parametrize('package', ['visualiser', 'visualiserTV'])
def test_visualizer_type_dates_and_filters(tmp_path, package):
    source = tmp_path / 'Schedule.xlsx'
    excel_exporter.create_excel_from_html_data(records(), source, sync_metadata=SYNC)
    processor = importlib.import_module(package + '.data_processor')
    df = processor.load_data(source)
    assert len(processor.filter_by_lesson_type(df, 'all')) == 6
    assert len(processor.filter_by_lesson_type(df, 'group')) == 1
    assert len(processor.filter_by_lesson_type(df, 'non-group')) == 5
    days, schedule = processor.process_schedule_data(df)
    assert days == ['Mo']
    rentals = [row for row in schedule['Mo'] if row['lesson_type'] == 'rental']
    assert len(rentals) == 4
    assert next(row for row in rentals if row['block_id'] == 'dated')['rental_dates'] == ['2026-09-28', '2026-10-05', '2026-10-12']
    assert next(row for row in rentals if row['block_id'] == '00017')['teacher'] == ''


@pytest.mark.parametrize('mutation', [
    {'rental_dates': ['2026-10-06']}, {'rental_dates_json': 'invalid'},
    {'rental_dates_json': '["2026-10-05"]'}, {'duration': 45}, {'duration': 60.5}, {'duration': 'invalid'},
    {'source_layer': 'base'}, {'lesson_type': 'group'}, {'block_metadata_json': '{"block_id":"other"}'},
])
def test_conflicting_export_data_rejected(tmp_path, mutation):
    with pytest.raises(excel_exporter.ExcelExportValidationError):
        excel_exporter.create_excel_from_html_data([{**records()[0], **mutation}], tmp_path / 'bad.xlsx')
    assert not (tmp_path / 'bad.xlsx').exists()


def test_duplicate_ids_rejected_at_all_python_boundaries(tmp_path):
    source = tmp_path / 'Schedule.xlsx'
    rows = records()
    rows[1]['block_id'] = rows[0]['block_id']
    with pytest.raises(excel_exporter.ExcelExportValidationError, match='Duplicate block_id'):
        excel_exporter.create_excel_from_html_data(rows, source)
    excel_exporter.create_excel_from_html_data(records(), source)
    book = openpyxl.load_workbook(source)
    book['Schedule'].cell(3, 12, book['Schedule'].cell(2, 12).value)
    book.save(source)
    book.close()
    with pytest.raises(ScheduleExchangeError, match='Duplicate block_id'):
        excel_parser.parse_schedule(source)
    planning = tmp_path / 'newpref.xlsx'
    write_planning_boundary(source, planning)
    with pytest.raises(ScheduleExchangeError, match='Duplicate block_id'):
        ScheduleReader(planning).read_excel()


@pytest.mark.parametrize('column,value', [(13, 'bad JSON'), (13, '["2026-10-06"]'), (12, ''), (9, 45), (9, 60.5), (6, 'unknown')])
def test_import_contradictions_fail_before_html_generation(tmp_path, column, value):
    source = tmp_path / 'bad.xlsx'
    excel_exporter.create_excel_from_html_data(records(), source, sync_metadata=SYNC)
    book = openpyxl.load_workbook(source)
    book['Schedule'].cell(2, column).value = value
    book.save(source)
    book.close()
    html_dir = tmp_path / 'html'
    html_dir.mkdir()
    html = html_dir / 'schedule.html'
    html.write_text('previous output', encoding='utf-8')
    with pytest.raises(SchedulePipelineError):
        SchedulePipeline().process_excel_to_outputs(str(source), {'html': str(html_dir)})
    assert html.read_text(encoding='utf-8') == 'previous output'


@pytest.mark.parametrize('section,column', [(0, 'B'), (0, 'X'), (99, 'B'), ('invalid', 'B')])
def test_planning_metadata_conflicts_are_not_silently_ignored(tmp_path, section, column):
    source, planning = tmp_path / 'source.xlsx', tmp_path / 'newpref.xlsx'
    excel_exporter.create_excel_from_html_data(records(), source, sync_metadata=SYNC)
    write_planning_boundary(source, planning)
    book = openpyxl.load_workbook(planning)
    book['__service_metadata'].append([section, column, 'rental', '', 'orphan-id', '[]'])
    book.save(planning)
    book.close()
    with pytest.raises(ScheduleExchangeError):
        ScheduleReader(planning).read_excel()


def test_legacy_trial_rental_dates_survive_compatibility_and_reexport(tmp_path):
    source = tmp_path / 'legacy-trial.xlsx'
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = 'Schedule'
    sheet.append(SCHEDULE_COLUMNS[:9] + ['lesson_type', 'trial_dates_json'])
    row = records()[0]
    sheet.append([row.get(key, '') for key in SCHEDULE_COLUMNS[:9]] + ['trial', '["2026-09-28","2026-10-05"]'])
    book.save(source)
    book.close()
    activity = excel_parser.parse_schedule(source)[1]
    assert activity['lesson_type'] == 'rental'
    assert activity['rental_dates'] == ['2026-09-28', '2026-10-05']
    assert activity['trial_dates'] == []
    assert excel_exporter.create_excel_from_html_data([activity], tmp_path / 'reexport.xlsx')


@pytest.mark.parametrize('package', ['visualiser', 'visualiserTV'])
def test_legacy_optimizer_pause_columns_are_not_lesson_type(tmp_path, package):
    source = tmp_path / 'legacy-optimized.xlsx'
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = 'Schedule'
    sheet.append(SCHEDULE_COLUMNS[:11])
    row = records()[-1]
    sheet.append([row.get(key, '1A' if key == 'group' else 0) for key in SCHEDULE_COLUMNS[:11]])
    book.save(source)
    book.close()
    assert excel_parser.parse_schedule(source)[1]['lesson_type'] == 'group'
    processor = importlib.import_module(package + '.data_processor')
    loaded = processor.load_data(source)
    assert len(processor.filter_by_lesson_type(loaded, 'group')) == 1
    assert loaded['pause_before'].tolist() == [0]


def test_legacy_unknown_origin_and_one_time_primary_identity(tmp_path, monkeypatch):
    source = tmp_path / 'legacy.xlsx'
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = 'Schedule'
    sheet.append(SCHEDULE_COLUMNS[:9])
    row = records()[0]
    sheet.append([row.get(key, row.get('students') if key == 'group' else '') for key in SCHEDULE_COLUMNS[:9]])
    book.save(source)
    book.close()
    first, second = excel_parser.parse_schedule(source), excel_parser.parse_schedule(source)
    assert first == second and first.sync_metadata is None
    assert first[1]['lesson_type'] == 'rental' and first[1]['block_id'] == ''
    blocks = collect_individual_blocks_from_buildings(build_schedule_structure(first))
    assert len(blocks) == 1 and blocks[0]['id'] == ''
    monkeypatch.setattr(integration, 'get_schedule_state_dir', lambda: str(tmp_path / 'state'))
    integration.check_excel_generation_origin(source)
    integration.reset_web_editor_state(blocks)
    path = tmp_path / 'state/individual_lessons.json'
    before = path.read_bytes()
    state = json.loads(before)
    assert state['blocks'][0]['id'] and state['blocks'][0]['id'] != '1'
    with pytest.raises(SchedulePipelineError, match='первичного импорта'):
        integration.reset_web_editor_state(blocks)
    with pytest.raises(SchedulePipelineError, match='первичного импорта'):
        integration.check_excel_generation_origin(source)
    assert path.read_bytes() == before


def test_present_empty_revisions_differ_from_missing_metadata(tmp_path):
    empty = dict(SYNC, source_base_revision='', source_individual_revision='')
    book = openpyxl.Workbook()
    assert read_sync_metadata(book) is None
    write_sync_metadata(book, empty)
    path = tmp_path / 'empty.xlsx'
    book.save(path)
    book.close()
    book = openpyxl.load_workbook(path)
    assert read_sync_metadata(book) == empty
    book.close()


def test_partial_snapshot_cannot_reset_state(tmp_path, monkeypatch):
    monkeypatch.setattr(integration, 'get_schedule_state_dir', lambda: str(tmp_path / 'state'))
    with pytest.raises(SchedulePipelineError, match='Частичный'):
        integration.reset_web_editor_state([], sync_metadata=dict(SYNC, snapshot_scope='partial'))
    assert not (tmp_path / 'state').exists()


def test_rental_class_only_html_is_removed():
    stripped, removed = strip_non_group_activity_blocks_from_html(
        "<div class='activity-block lesson-type-rental'>Rental</div><div class='activity-block lesson-type-group'>Lesson</div>")
    assert removed == 1 and 'Rental' not in stripped and 'Lesson' in stripped


def test_export_api_preserves_supplied_revisions_and_legacy_origin(isolated_editor, tmp_path, monkeypatch):
    routes, ind_path, base_path = isolated_editor
    monkeypatch.setattr(routes, 'EXCEL_EXPORTS_DIR', str(tmp_path / 'exports'))
    # Phase 4 labels only a payload whose revisions and managed IDs match the server.
    base_path.write_text(json.dumps({'published_at': 'base-old', 'blocks': []}), encoding='utf-8')
    ind_path.write_text(json.dumps({'last_modified': None, 'blocks': [
        {'id': row['block_id']} for row in records() if row['block_id']]}), encoding='utf-8')
    with routes.app.test_client() as client:
        login(client, 'admin')
        for sync in (SYNC, None):
            form = {'schedule_data': json.dumps(records())}
            if sync is not None:
                form.update(schedule_sync=json.dumps(sync), schedule_html_revision='')
            response = client.post('/export_to_excel', data=form)
            assert response.status_code == 200
            from io import BytesIO
            book = openpyxl.load_workbook(BytesIO(response.data))
            assert read_sync_metadata(book) == sync
            book.close()
        bad = client.post('/export_to_excel', data={'schedule_data': json.dumps(records()), 'schedule_sync': '{"format_version":1}'})
        assert bad.status_code == 400


def test_vba_transfer_and_both_import_paths_are_static():
    macro = (ROOT / 'gear_xls/Modul1.bas').read_text(encoding='utf-8')
    assert macro.isascii(), 'VBComponents.Import must not depend on the Windows ANSI code page'
    for column, key in enumerate(['section_index', 'column_letter', 'lesson_type', 'trial_dates_json', *RECORD_COLUMNS], 1):
        assert f'wsMetadata.Cells(1, {column}).Value = "{key}"' in macro
    # Worksheet.Copy of the very hidden sync sheet fails in real Excel; cells are transferred instead.
    assert 'wsSync.Copy' not in macro and 'Set wsSyncTarget = wbTarget.Worksheets.Add(After:=wsMetadata)' in macro
    assert 'wsSyncTarget.Range(syncCell.Address).NumberFormat = syncCell.NumberFormat' in macro
    assert 'wsSyncTarget.Visible = xlSheetVeryHidden' in macro
    assert 'wsSync = wbSource.Worksheets("__schedule_sync")' in macro
    assert 'targetRow = targetRow + 14' in macro
    assert 'wsTarget.Cells(targetRow + 11, targetColumn).Value = wsSource.Cells(i, 8).Value' in macro
    for filename in ('gear_xls/convert_to_xlsm.py', 'gui_services/app_actions.py'):
        code = (ROOT / filename).read_text(encoding='utf-8-sig')
        assert '"Modul1.bas"' in code and 'VBComponents.Import(module_path)' in code
        assert 'FileFormat=52' in code and 'CreateSchedulePlanning' in code
        ast.parse(code)


def test_solver_carries_identity_dates_and_provenance_when_available(tmp_path):
    pytest.importorskip('ortools.sat.python.cp_model', reason='Real OR-Tools unavailable; no solver substitute')
    from scheduler_base import ScheduleOptimizer
    source, planning = tmp_path / 'source.xlsx', tmp_path / 'newpref.xlsx'
    excel_exporter.create_excel_from_html_data(records(), source, sync_metadata=SYNC)
    write_planning_boundary(source, planning)
    reader = ScheduleReader(planning)
    optimizer = ScheduleOptimizer(reader.read_excel(), time_interval=5, sync_metadata=reader.sync_metadata)
    assert optimizer.solve(time_limit_seconds=5)
    assert optimizer.sync_metadata == SYNC
    output = tmp_path / 'actual_solution.xlsx'
    assert export_to_excel(optimizer, output)
    assert_activities(excel_parser.parse_schedule(output))
