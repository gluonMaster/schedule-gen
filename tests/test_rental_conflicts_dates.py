"""Phase 2: shared Python/JS cases and optional real 2–3-event solver models."""
import json
import sys
from datetime import date
from pathlib import Path
from types import SimpleNamespace

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from reader import ScheduleClass
from rental_conflicts import fixed_conflict_type, rental_calendar_overlap
from conflict_detector import check_potential_conflicts
from gear_xls import rooms_report, rooms_routes, state_manager

TABLE = json.loads((ROOT / "tests/fixtures/rental_conflicts_cases.fixture").read_text(encoding="utf-8"))
TODAY = date.fromisoformat(TABLE["calculation_date"])
CASES = TABLE["cases"]


def record(updates=None):
    return {**TABLE["defaults"], **(updates or {})}


def event(data=None, **updates):
    data = record({**(data or {}), **updates})
    return ScheduleClass(
        subject=data["subject"], group=data["students"], teacher=data["teacher"],
        main_room=data["room"], alternative_rooms=data.get("alternative_rooms", []),
        building=data["building"], duration=data["duration"], day=data["day"],
        start_time=data["start_time"],
        end_time=data["end_time"] if data["lesson_type"] == "rental" or data.get("window") else None,
        lesson_type=data["lesson_type"], rental_dates=data["rental_dates"],
        trial_dates=data["trial_dates"], block_id=data.get("block_id", ""),
    )


@pytest.mark.parametrize("case", CASES, ids=lambda case: case["name"])
def test_fixed_conflicts_and_rental_diagnostics(case):
    first, second = event(case.get("first")), event(case.get("second"))
    assert fixed_conflict_type(first, second, TODAY) == case["conflict"]
    diagnostics = check_potential_conflicts(SimpleNamespace(classes=[first, second], calculation_date=TODAY))
    if first.is_rental or second.is_rental:
        assert bool(diagnostics) == bool(case["conflict"])
    if first.is_rental:
        assert first.rental_dates == record(case.get("first"))["rental_dates"]


def test_rental_preserves_display_fields_but_has_no_teaching_resources():
    booking = event(block_id="booking-42", rental_dates=["2026-09-28"], alternative_rooms=["1.02"])
    assert (booking.teacher, booking.group, booking.block_id) == ("Renter", "Organisation", "booking-42")
    assert booking.resource_teacher == "" and booking.get_groups() == []
    assert booking.possible_rooms == ["1.01"] and booking.has_fixed_room and booking.has_fixed_time
    assert not booking.has_time_window
    assert not rental_calendar_overlap(booking, event(), TODAY)
    assert booking.rental_dates == ["2026-09-28"]


def test_variable_day_is_not_filtered_out_before_day_vars():
    booking = event(rental_dates=["2026-10-05"])
    lesson = event(lesson_type="group", day="", teacher="Teacher", students="2A")
    assert rental_calendar_overlap(booking, lesson, TODAY)
    assert check_potential_conflicts(SimpleNamespace(classes=[booking, lesson], calculation_date=TODAY)) == []
    lesson.day = "Di"
    assert not rental_calendar_overlap(booking, lesson, TODAY)


def test_window_heuristics_leave_rentals_out_of_teaching_and_window_terms():
    from timewindow_adapter import apply_timewindow_improvements, add_objective_weights_for_timewindows
    bookings = [event(rental_dates=["2026-10-05"]), event(rental_dates=["2026-10-12"])]
    # These pure exclusions must not need or add any model constraints.
    optimizer = SimpleNamespace(classes=bookings, start_vars={0: 8, 1: 8}, model=None)
    assert apply_timewindow_improvements(optimizer)
    assert add_objective_weights_for_timewindows(optimizer) == []
    assert optimizer.prefer_late_start == set()
    assert all(c.has_fixed_time and not c.has_time_window and c.teacher == "Renter" for c in bookings)


def test_linked_constraints_do_not_treat_rentals_as_sequential_lessons():
    from linked_constraints import add_linked_constraints, build_linked_chains
    first, second = event(), event(room="1.02")
    first.linked_classes = [second]
    classes = [first, second]
    optimizer = SimpleNamespace(classes=classes, _find_class_index=classes.index, model=None)
    add_linked_constraints(optimizer)
    assert optimizer.linked_chains == []
    lesson = event(lesson_type="group")
    lesson.linked_classes = [first]
    classes.append(lesson)
    add_linked_constraints(optimizer)
    assert all(0 not in chain and 1 not in chain for chain in optimizer.linked_chains)
    next_lesson = event(lesson_type="group", subject="Next", start_time="11:00", end_time="12:00")
    lesson.linked_classes = [next_lesson]
    classes.append(next_lesson)
    build_linked_chains(optimizer)
    assert optimizer.linked_chains == [[2, 3]]


def test_reader_planning_links_only_teaching_columns(tmp_path):
    from openpyxl import Workbook
    from reader import ScheduleReader
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Plannung"
    metadata = workbook.create_sheet("__service_metadata")
    metadata.append(["section_index", "column_letter", "lesson_type", "trial_dates_json"])
    for col, subject, kind in [(2, "Math", "group"), (3, "Vermietung", "rental"), (4, "Art", "group")]:
        teacher = "Renter" if kind == "rental" else "Teacher"
        for offset, value in enumerate([subject, "2A", teacher, "1.01", None, None, None,
                                        "Villa", 60, "Mo", "10:00", None, 0, 0]):
            sheet.cell(row=2 + offset, column=col, value=value)
        metadata.append([0, chr(64 + col), kind, ""])
    source = tmp_path / "planning-links.xlsx"
    workbook.save(source)
    reader = ScheduleReader(str(source))
    first, booking, last = reader.read_excel()
    assert first.linked_classes == [last]
    assert booking.linked_classes == [] and booking.previous_class is None and booking.next_class is None
    assert first.next_class == "Art" and last.previous_class == "Math"
    assert booking.teacher == "Renter" and booking.group == "2A"
    assert reader.teachers == {"Teacher"}


@pytest.fixture
def real_optimizer():
    pytest.importorskip("ortools.sat.python.cp_model", reason="Real OR-Tools is unavailable; solver not verified")
    from scheduler_base import ScheduleOptimizer
    return lambda classes: ScheduleOptimizer(classes, time_interval=15, calculation_date=TODAY)


@pytest.mark.parametrize("case", CASES, ids=lambda case: case["name"])
def test_real_solver_shared_cases(real_optimizer, case):
    classes = [event(case.get("first")), event(case.get("second"))]
    optimizer = real_optimizer(classes)
    success = optimizer.solve(time_limit_seconds=1)
    assert success == (case["conflict"] is None), optimizer.last_status_name
    assert optimizer.last_status_name in ({"OPTIMAL", "FEASIBLE"} if success else {"INFEASIBLE"})
    if success:
        for source, output in zip(classes, optimizer.solution):
            if source.is_rental:
                assert output["start_time"] == source.start_time
                assert output["duration"] == source.duration and output["room"] == source.main_room
                assert output["lesson_type"] == "rental"
                assert json.loads(output["rental_dates_json"]) == source.rental_dates


@pytest.mark.parametrize("dates", [[], ["2026-10-05"], ["2026-09-28"]])
def test_real_solver_free_alternative_and_fixed_rental(real_optimizer, dates):
    booking = event(rental_dates=dates, block_id="stable-booking", alternative_rooms=["1.03"])
    lesson = event(lesson_type="group", window=True, alternative_rooms=["1.02"])
    optimizer = real_optimizer([booking, lesson])
    assert optimizer.solve(time_limit_seconds=1), optimizer.last_status_name
    result = optimizer.solution
    assert result[0]["room"] == "1.01" and result[0]["start_time"] == "10:00"
    assert result[0]["end_time"] == "11:00" and result[0]["block_id"] == "stable-booking"
    assert result[1]["start_time"] == "10:00" and result[1]["duration"] == 60
    if not dates or dates[0] >= TODAY.isoformat():
        assert result[1]["room"] == "1.02"


@pytest.mark.parametrize("chosen_day,feasible", [("Mo", False), ("Di", True)])
def test_real_solver_variable_day_obeys_rental(real_optimizer, chosen_day, feasible):
    booking = event(rental_dates=["2026-10-05"])
    lesson = event(lesson_type="group", day="", teacher="Teacher", students="2A", window=True)
    anchor = event(lesson_type="group", day="Di", room="1.03", teacher="Other", students="3A")
    optimizer = real_optimizer([booking, lesson, anchor])
    optimizer.build_model()
    optimizer.model.Add(optimizer.day_vars[1] == optimizer.day_indices[chosen_day])
    assert optimizer.solve(time_limit_seconds=1) == feasible, optimizer.last_status_name


def test_real_solver_rental_does_not_create_a_teaching_link(real_optimizer):
    first, second = event(), event(room="1.02")
    first.linked_classes = [second]
    optimizer = real_optimizer([first, second])
    assert optimizer.solve(time_limit_seconds=1), optimizer.last_status_name
    assert all(item["start_time"] == "10:00" for item in optimizer.solution)


def test_rooms_report_passes_rental_dates_without_editing_sources(tmp_path, monkeypatch):
    booking = record({"id": "room-booking", "rental_dates": ["2026-09-28", "2026-10-05"]})
    state_path = tmp_path / "individual_lessons.json"
    state_path.write_text(json.dumps({"blocks": [booking]}), encoding="utf-8")
    monkeypatch.setattr(state_manager, "INDIVIDUAL_LESSONS_PATH", str(state_path))
    monkeypatch.setattr(state_manager, "INDIVIDUAL_LOCK_PATH", str(tmp_path / "individual.lock"))
    monkeypatch.setattr(state_manager, "_today_local_date", lambda: TODAY)
    monkeypatch.setattr(rooms_report, "_load_base_blocks", lambda: [])
    monkeypatch.setattr(rooms_report, "_load_configured_rooms", lambda: {"Villa": ["1.01"]})
    before = state_path.read_bytes()
    result = rooms_report.compute_availability()
    slot = result["buildings"]["Villa"]["days"]["Mo"]["1.01"][0]
    assert slot["lesson_type"] == "rental" and slot["rental_dates"] == booking["rental_dates"]
    assert slot["teacher"] == "Renter" and slot["students"] == "Organisation"
    assert state_path.read_bytes() == before
    assert "rooms_report.js?v=20261001_rental2" in rooms_routes.ROOMS_PAGE_TEMPLATE
